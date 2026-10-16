"""Plantilla Excel estándar para que cada área cargue proyecciones y pedidos.

Hoja `datos` con columnas fijas y validaciones (listas desplegables, fecha, kg ≥ 0).
Las listas salen del catálogo (backend/config/catalog.yaml) y viven en una hoja oculta
`listas`, para no depender del límite de 255 caracteres de las listas en línea.
"""

from collections.abc import Iterable, Sequence
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.core.catalog import Catalog

DATA_SHEET = "datos"
LISTS_SHEET = "listas"
HELP_SHEET = "instrucciones"

# (encabezado, ancho, formato de número)
COLUMNS: list[tuple[str, int, str | None]] = [
    ("fecha", 12, "yyyy-mm-dd"),
    ("canal", 16, None),
    ("cliente", 22, None),
    ("producto", 16, None),
    ("cantidad_kg", 13, "0.0"),
    ("tipo", 12, None),
    ("responsable", 18, None),
]
HEADERS = [c[0] for c in COLUMNS]
DEFAULT_ROWS = 2000

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def file_name(channel: str, responsible: str, day: date) -> str:
    """Nombre de archivo según la convención `{canal}_{responsable}_{AAAAMMDD}.xlsx`."""
    return f"{channel}_{responsible}_{day:%Y%m%d}.xlsx"


def _col(header: str) -> str:
    return get_column_letter(HEADERS.index(header) + 1)


def _list_validation(
    lists: Worksheet, col: int, values: Sequence[str], title: str
) -> DataValidation:
    letter = get_column_letter(col)
    lists.cell(row=1, column=col, value=title).font = Font(bold=True)
    for i, value in enumerate(values, start=2):
        lists.cell(row=i, column=col, value=value)
    ref = f"={LISTS_SHEET}!${letter}$2:${letter}${len(values) + 1}"
    dv = DataValidation(type="list", formula1=ref, allow_blank=True)
    dv.error = f"Elige un valor de la lista de {title}."
    dv.prompt = f"Elige {title} de la lista."
    return dv


def _set_messages(dv: DataValidation, title: str) -> DataValidation:
    dv.errorTitle = "Valor no válido"
    dv.promptTitle = title
    dv.errorStyle = "stop"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    return dv


def build_template(catalog: Catalog, rows: int = DEFAULT_ROWS) -> Workbook:
    """Crea el libro de la plantilla con validaciones para `rows` filas de datos."""
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    lists = wb.create_sheet(LISTS_SHEET)
    help_ws = wb.create_sheet(HELP_SHEET)

    for i, (header, width, fmt) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if fmt:
            for r in range(2, rows + 2):
                ws.cell(row=r, column=i).number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    last = rows + 1
    validations: dict[str, DataValidation] = {
        "canal": _list_validation(lists, 1, catalog.channels, "canal"),
        "producto": _list_validation(lists, 2, catalog.products, "producto"),
        "tipo": _list_validation(lists, 3, catalog.record_types, "tipo"),
    }
    dv_date = DataValidation(
        type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2100,12,31)"
    )
    dv_date.error = "Ingresa una fecha válida (AAAA-MM-DD)."
    dv_date.prompt = "Fecha del pedido o de la proyección, formato AAAA-MM-DD."
    validations["fecha"] = dv_date
    dv_qty = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
    dv_qty.error = "La cantidad debe ser un número mayor o igual a 0, en kg."
    dv_qty.prompt = "Cantidad en kilogramos (número ≥ 0)."
    validations["cantidad_kg"] = dv_qty

    for header, dv in validations.items():
        _set_messages(dv, header)
        dv.add(f"{_col(header)}2:{_col(header)}{last}")
        ws.add_data_validation(dv)
    lists.sheet_state = "hidden"

    _write_help(help_ws, catalog)
    return wb


def _write_help(ws: Worksheet, catalog: Catalog) -> None:
    lines = [
        ("Plantilla de carga — Avícola 2", True),
        ("", False),
        ("Completa la hoja «datos», una fila por cliente, producto y fecha.", False),
        ("No cambies los nombres ni el orden de las columnas.", False),
        ("", False),
        ("Columnas", True),
        ("fecha: AAAA-MM-DD (obligatoria).", False),
        (f"canal: {', '.join(catalog.channels)} (obligatoria).", False),
        ("cliente: nombre o código del cliente (opcional).", False),
        (f"producto: {', '.join(catalog.products)} (obligatoria).", False),
        ("cantidad_kg: número ≥ 0, en kilogramos (obligatoria).", False),
        (f"tipo: {' / '.join(catalog.record_types)} (obligatoria).", False),
        ("responsable: quién registra los datos (opcional).", False),
        ("", False),
        ("Nombre del archivo", True),
        (
            "{canal}_{responsable}_{AAAAMMDD}.xlsx — por ejemplo: Horeca_comercial1_20261015.xlsx",
            False,
        ),
    ]
    ws.column_dimensions["A"].width = 95
    for i, (text, bold) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text).font = Font(bold=bold, size=13 if i == 1 else 11)


def fill_rows(wb: Workbook, rows: Iterable[Sequence[Any]]) -> int:
    """Escribe filas (en el orden de HEADERS) en la hoja `datos`. Devuelve cuántas escribió."""
    ws = wb[DATA_SHEET]
    n = 0
    for n, row in enumerate(rows, start=1):
        for j, value in enumerate(row, start=1):
            ws.cell(row=n + 1, column=j, value=value)
    return n
