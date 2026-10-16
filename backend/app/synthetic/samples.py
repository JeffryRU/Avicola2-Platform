"""Excel de ejemplo con formatos distintos, para probar la ingesta y la normalización con IA.

Cada archivo representa un caso que la ingesta debe resolver (ver data/samples/README.md):
    1. plantilla estándar                  → se ingiere directo
    2. columnas con otros nombres          → se resuelve con los alias del YAML
    3. resumen semanal sin producto        → granularidad semanal, producto TOTAL, canal del nombre
    4. columnas desconocidas y formato local → necesita mapeo propuesto por el LLM
    5. plantilla con filas inválidas       → la ingesta reporta errores por fila
"""

from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from app.core.catalog import Catalog
from app.ingest.template import build_template, file_name, fill_rows
from app.synthetic.generator import SyntheticDataset


@dataclass(frozen=True)
class Sample:
    path: Path
    case: str


def _last_days(df: pd.DataFrame, days: int) -> pd.DataFrame:
    cutoff = df["date"].max() - pd.Timedelta(days=days - 1)
    return df[df["date"] >= cutoff]


def _template_ok(ds: SyntheticDataset, catalog: Catalog, out: Path) -> Sample:
    rows = _last_days(ds.orders[ds.orders["channel"] == "Horeca"], 3)
    wb = build_template(catalog)
    fill_rows(
        wb,
        (
            (r.date.date(), r.channel, r.customer, r.product, r.ordered_kg, "pedido", "comercial1")
            for r in rows.itertuples()
        ),
    )
    path = out / file_name("Horeca", "comercial1", ds.config.end)
    wb.save(path)
    return Sample(path, "plantilla estándar")


def _aliases(ds: SyntheticDataset, out: Path) -> Sample:
    rows = _last_days(ds.orders[ds.orders["channel"] == "Supermercados"], 3)
    df = pd.DataFrame(
        {
            "Fecha Pedido": rows["date"].dt.date,
            "Canal": rows["channel"],
            "Tienda": rows["customer"],
            "Item": rows["product"],
            "Kilos": rows["ordered_kg"],
            "Tipo": "pedido",
        }
    )
    path = out / file_name("Supermercados", "comercial2", ds.config.end)
    df.to_excel(path, sheet_name="Pedidos", index=False)
    return Sample(path, "columnas con alias conocidos")


def _weekly_total(ds: SyntheticDataset, out: Path) -> Sample:
    tiendas = ds.orders[ds.orders["channel"] == "Tiendas"]
    weekly = tiendas.groupby(pd.Grouper(key="date", freq="W-MON", label="left", closed="left"))[
        "ordered_kg"
    ].sum()
    # Solo semanas completas (lunes a domingo) dentro del histórico.
    complete = weekly.index + pd.Timedelta(days=6) <= pd.Timestamp(ds.config.end)
    weekly = weekly[complete].tail(8)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Reporte semanal — Tiendas"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Semana", "Total Kg"])
    for week, kg in weekly.items():
        ws.append([week.date(), round(float(kg), 1)])
    ws.append([])
    ws.append(["Fuente: consolidado manual del área"])
    path = out / file_name("Tiendas", "jefe-tiendas", ds.config.end)
    wb.save(path)
    return Sample(path, "resumen semanal sin producto ni cliente")


def _needs_llm(ds: SyntheticDataset, out: Path, rng: np.random.Generator) -> Sample:
    rows = _last_days(ds.orders[ds.orders["channel"] == "Horeca"], 2)
    projected = rows["ordered_kg"].to_numpy() * rng.lognormal(0.05, 0.1, len(rows))
    df = pd.DataFrame(
        {
            "Día": rows["date"].dt.strftime("%d/%m/%Y").to_numpy(),
            "Local": rows["customer"].to_numpy(),
            "Descripción": rows["product"].str.upper().to_numpy(),
            # Formato local: punto de miles y coma decimal, como texto.
            "Peso neto (kg)": [
                f"{v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".") for v in projected
            ],
            "Observaciones": "proyección enviada por WhatsApp",
        }
    )
    path = out / file_name("Horeca", "horeca-wsp", ds.config.end)
    df.to_excel(path, sheet_name="Hoja1", index=False)
    return Sample(path, "columnas desconocidas y formato local (requiere LLM)")


def _with_errors(ds: SyntheticDataset, catalog: Catalog, out: Path) -> Sample:
    rows = _last_days(ds.orders[ds.orders["channel"] == "Supermercados"], 1)
    good = [
        (r.date.date(), r.channel, r.customer, r.product, r.ordered_kg, "pedido", "comercial3")
        for r in rows.itertuples()
    ]
    day = ds.config.end
    bad = [
        (day, "Supermercados", "SUP-01", "Pechuga", -25.0, "pedido", "comercial3"),  # kg negativo
        (None, "Supermercados", "SUP-02", "Pierna", 310.0, "pedido", "comercial3"),  # sin fecha
        (
            day,
            "Supermercados",
            "SUP-03",
            "Pavo",
            120.0,
            "pedido",
            "comercial3",
        ),  # producto fuera del catálogo
        (day, "Supermercado", "SUP-01", "Alas", 80.0, "pedido", "comercial3"),  # canal mal escrito
        (
            day,
            "Supermercados",
            "SUP-02",
            "Pechuga",
            "n/a",
            "pedido",
            "comercial3",
        ),  # kg no numérico
        (
            day,
            "Supermercados",
            "SUP-03",
            "Pollo entero",
            500.0,
            "venta",
            "comercial3",
        ),  # tipo inválido
    ]
    wb = build_template(catalog)
    fill_rows(
        wb, good + bad
    )  # openpyxl no aplica validaciones al escribir: simula una carga sin Excel
    path = out / file_name("Supermercados", "comercial3", day)
    wb.save(path)
    return Sample(path, f"plantilla con {len(bad)} filas inválidas")


def write_samples(ds: SyntheticDataset, catalog: Catalog, out_dir: Path) -> list[Sample]:
    """Genera los 5 Excel de ejemplo a partir de los últimos días del dataset."""
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(ds.config.seed + 1)
    return [
        _template_ok(ds, catalog, out_dir),
        _aliases(ds, out_dir),
        _weekly_total(ds, out_dir),
        _needs_llm(ds, out_dir, rng),
        _with_errors(ds, catalog, out_dir),
    ]
