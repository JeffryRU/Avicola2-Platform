from datetime import date

import pandas as pd
from openpyxl import load_workbook

from app.core.catalog import get_catalog
from app.ingest.template import (
    DATA_SHEET,
    HEADERS,
    LISTS_SHEET,
    build_template,
    file_name,
    fill_rows,
)
from app.synthetic.generator import SyntheticConfig, generate
from app.synthetic.samples import write_samples


def test_template_structure_and_validations(tmp_path) -> None:
    catalog = get_catalog()
    path = tmp_path / "plantilla.xlsx"
    build_template(catalog, rows=50).save(path)

    wb = load_workbook(path)
    ws = wb[DATA_SHEET]
    assert [c.value for c in ws[1]] == HEADERS
    assert wb[LISTS_SHEET].sheet_state == "hidden"

    by_range = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    assert by_range["B2:B51"].formula1 == "=listas!$A$2:$A$4"  # canal
    assert by_range["D2:D51"].type == "list"  # producto
    assert by_range["F2:F51"].type == "list"  # tipo
    assert by_range["A2:A51"].type == "date"
    assert by_range["E2:E51"].type == "decimal"
    assert by_range["E2:E51"].formula1 == "0"
    assert all(dv.showErrorMessage for dv in by_range.values())

    products = [c.value for c in wb[LISTS_SHEET]["B"][1:] if c.value]
    assert products == catalog.products


def test_fill_rows_round_trip(tmp_path) -> None:
    wb = build_template(get_catalog(), rows=10)
    n = fill_rows(wb, [(date(2026, 10, 1), "Horeca", "HOR-01", "Pechuga", 12.5, "pedido", "x")])
    path = tmp_path / "f.xlsx"
    wb.save(path)
    df = pd.read_excel(path, sheet_name=DATA_SHEET).dropna(how="all")
    assert n == 1
    assert df.iloc[0]["producto"] == "Pechuga"
    assert df.iloc[0]["cantidad_kg"] == 12.5


def test_file_name_convention() -> None:
    assert (
        file_name("Horeca", "comercial1", date(2026, 10, 15)) == "Horeca_comercial1_20261015.xlsx"
    )


def test_samples_cover_all_cases(tmp_path) -> None:
    ds = generate(SyntheticConfig(end=date(2026, 10, 14), days=60))
    samples = write_samples(ds, get_catalog(), tmp_path)
    assert len(samples) == 5
    assert all(s.path.exists() for s in samples)

    weekly = next(s for s in samples if s.path.name.startswith("Tiendas"))
    df = pd.read_excel(weekly.path, skiprows=2).dropna()
    weeks = pd.to_datetime(df["Semana"])
    assert (weeks.dt.dayofweek == 0).all()
    assert (weeks + pd.Timedelta(days=6) <= pd.Timestamp("2026-10-14")).all()
